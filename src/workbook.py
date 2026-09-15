"""PM-facing workbook: the nowcast, its track record and the panel trends, formatted for Excel.

Anything derived from other cells is a live Excel formula (panel change, forecast, model versus
Street, error columns, averages and the sentences that quote them). Only pipeline outputs are
written as values: reported comps, panel YoY, consensus, fitted slopes, per-day errors and trends.
"""

import pandas as pd
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TITLES = {
    "ANTHROPOLOGIE GROUP (US)": "Anthropologie",
    "FREE PEOPLE (US)": "Free People",
    "URBAN OUTFITTERS (US)": "Urban Outfitters",
    "Total": "URBN total",
}
WINDOWS = {
    "daily_spend_yoy": "Daily YoY",
    "t7d_spend_yoy": "Trailing 7-day YoY",
    "qtd_spend_yoy": "QTD YoY",
}
# One colour per entity, shared with the charts in report.py.
COLORS = {
    "model": "#2a78d6",
    "reported": "#eb6834",
    "naive": "#1baf7a",
    "consensus": "#eda100",
    "raw_panel": "#e87ba4",
    "header": "#184f95",
}
HEADER = PatternFill("solid", fgColor=COLORS["header"][1:])
HEADER_FONT = Font(bold=True, color="FFFFFF")
HIGHLIGHT = PatternFill("solid", fgColor="E3EEFB")
NEGATIVE, POSITIVE, WHITE = COLORS["reported"][1:], COLORS["naive"][1:], "FFFFFF"
PERCENT, POINTS, RATIO, DATE = "0.0%", "+0.0;-0.0;0.0", "0.00", "yyyy-mm-dd"
TEXT = {"KPI", "Quarter", "Day in quarter", "Fiscal quarter", "Topic", "Note"}
IMAGE_SIZE, IMAGE_ROWS = (900, 600), 32  # half the PNG size, and the rows it covers
SUMMARY_ROWS = range(
    4, 4 + len(TITLES)
)  # title in row 1, header in row 3, one row per KPI
DIRECTION = "Read the forecast as a direction versus consensus, not a point estimate."


def number_format(label):
    """Excel format from the header: points for (pp), the slope as a ratio, dates, text, else a rate."""
    if "(pp)" in label:
        return POINTS
    if "slope" in label.lower():
        return RATIO
    if label == "Date":
        return DATE
    return None if label in TEXT else PERCENT


def diverging(limit):
    """Orange at -limit, white at zero, green at +limit."""
    return ColorScaleRule(
        start_type="num",
        start_value=-limit,
        start_color=NEGATIVE,
        mid_type="num",
        mid_value=0,
        mid_color=WHITE,
        end_type="num",
        end_value=limit,
        end_color=POSITIVE,
    )


def sequential():
    """White at zero to orange at the column maximum, for error sizes."""
    return ColorScaleRule(
        start_type="num",
        start_value=0,
        start_color=WHITE,
        end_type="max",
        end_color=NEGATIVE,
    )


def context(latest):
    """As-of date, day-in-quarter, in-progress quarter and the latest reported quarter."""
    row = latest.iloc[0]
    return row["date"], int(row["diq"]), row["quarter_name"], row["ref_quarter"]


def style(ws, header_rows, formats, freeze_column=2):
    """Blue header band, widths and number formats by column, panes frozen below the header."""
    first_data = header_rows.stop
    for row in header_rows:
        ws.row_dimensions[row].height = 34
        for cell in ws[row]:
            cell.fill, cell.font = HEADER, HEADER_FONT
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
    for index, fmt in enumerate(formats, start=1):
        letter = get_column_letter(index)
        label = str(ws.cell(row=header_rows[-1], column=index).value or "")
        ws.column_dimensions[letter].width = min(max(len(label) + 2, 12), 24)
        for cell in ws[letter][first_data - 1 :] if fmt else ():
            cell.number_format = fmt
    ws.freeze_panes = ws.cell(row=first_data, column=freeze_column)


def shade(ws, header_row, match, rule):
    """Apply one colour-scale rule to the data of every column whose header matches."""
    ranges = [
        f"{get_column_letter(cell.column)}{header_row + 1}:{get_column_letter(cell.column)}{ws.max_row}"
        for cell in ws[header_row]
        if match(str(cell.value))
    ]
    ws.conditional_formatting.add(" ".join(ranges), rule)


def add_image(ws, path, anchor):
    image = Image(str(path))
    image.width, image.height = IMAGE_SIZE
    ws.add_image(image, anchor)


def kpi_table(frame, index, labels):
    """Wide table with one column group per KPI and one column per label, in TITLES order."""
    values = [column for column in labels if column in frame]
    wide = frame.pivot_table(index=index, columns="segment_name", values=values)
    wide.columns = wide.columns.swaplevel()
    wide = wide.reindex(
        columns=pd.MultiIndex.from_product([list(TITLES), list(labels)])
    )
    return wide.rename(columns=TITLES, level=0).rename(columns=labels, level=1)


def write_grouped(writer, name, table, key_label, freeze_column=2):
    """A table with KPI group headers over measure headers; the first column is the row key."""
    table.to_excel(writer, sheet_name=name)
    ws = writer.sheets[name]
    # pandas leaves a blank row for the index name under a two-level header.
    ws.delete_rows(3)
    ws.cell(row=2, column=1, value=key_label)
    formats = [
        number_format(key_label),
        *(number_format(label) for _, label in table.columns),
    ]
    style(ws, header_rows=range(1, 3), formats=formats, freeze_column=freeze_column)
    return ws


def span(quarters, title, measure):
    """Excel range of one KPI's column on the By quarter sheet, e.g. 'By quarter'!B3:B15."""
    letter, rows = quarters["letters"][title, measure], quarters["rows"]
    return f"'By quarter'!{letter}{rows.start}:{letter}{rows[-1]}"


def mae_formula(quarters, title, measure):
    """MAE in pp of a By quarter column against Reported, counting only quarters with a comp."""
    reported, column = span(quarters, title, "Reported"), span(quarters, title, measure)
    return (
        f'=100*SUMPRODUCT(ABS({column}-{reported})*({reported}<>""))/COUNT({reported})'
    )


def write_by_quarter(writer, backtest, diq, quarter):
    """PM and analyst view: each quarter's reported comp against the day-N forecast and the naive baseline."""
    measures = {
        "actual": "Reported",
        "forecast": f"Model (day {diq})",
        "ref_actual": "Naive",
        "error_pp": "Error (pp)",
    }
    rows = backtest[backtest["diq"].eq(diq)].dropna(subset=["forecast"])
    wide = kpi_table(rows, "quarter_name", measures)
    letters = {
        column: get_column_letter(index + 2)
        for index, column in enumerate(wide.columns)
    }
    data_rows = range(3, 3 + len(wide))
    for title in TITLES.values():
        reported, model = (
            letters[title, "Reported"],
            letters[title, measures["forecast"]],
        )
        # Live formula: 100 x (model - reported), blank until the comp is reported.
        wide[title, measures["error_pp"]] = [
            f'=IF({reported}{r}="","",100*({model}{r}-{reported}{r}))'
            for r in data_rows
        ]
    wide = wide.rename(index={quarter: f"{quarter} (in progress)"})
    ws = write_grouped(writer, "By quarter", wide, key_label="Quarter")
    shade(ws, 2, lambda label: "(pp)" in label, diverging(5))
    return {"letters": letters, "rows": data_rows}


def write_by_day(writer, errors, diq):
    """Analyst view: MAE of a forecast made on each day of the quarter, model against naive."""
    labels = {"model": "Model MAE (pp)", "naive": "Naive MAE (pp)"}
    by_method = errors.pivot_table(
        index=["diq", "segment_name"], columns="method", values="mae_pp"
    )
    mae = kpi_table(by_method.reset_index(), "diq", labels)
    data_rows = range(3, 3 + len(mae))
    for label in labels.values():
        # Live formula: the plain average of the four KPI columns for this method.
        kpi_columns = [
            get_column_letter(index + 2)
            for index, column in enumerate(mae.columns)
            if column[1] == label
        ]
        mae["Average of 4 KPIs", label] = [
            "=AVERAGE(" + ",".join(f"{letter}{r}" for letter in kpi_columns) + ")"
            for r in data_rows
        ]
    ws = write_grouped(writer, "By day", mae, key_label="Day in quarter")
    shade(ws, 2, lambda label: "(pp)" in label, sequential())
    for row in ws.iter_rows(min_row=3):
        if row[0].value == diq:
            for cell in row:
                cell.fill = HIGHLIGHT


def write_trends(writer, daily, chart, months=13):
    """Analyst and PM view: daily, trailing-7-day and QTD panel YoY by KPI, newest first."""
    recent = daily[daily["date"].ge(daily["date"].max() - pd.DateOffset(months=months))]
    wide = kpi_table(recent, "date", WINDOWS)
    keys = recent.drop_duplicates("date").set_index("date")[["quarter_name", "diq"]]
    keys.columns = pd.MultiIndex.from_tuples(
        [("", "Fiscal quarter"), ("", "Day in quarter")]
    )
    table = pd.concat([keys, wide], axis=1).sort_index(ascending=False)
    ws = write_grouped(writer, "Panel trends", table, key_label="Date", freeze_column=4)
    shade(ws, 2, lambda label: label.endswith("YoY"), diverging(0.2))
    add_image(ws, chart, f"A{ws.max_row + 2}")  # below the table, clear of the panes


def write_summary(writer, latest, charts, quarters):
    """PM view: where the model sits versus the last comp and the Street, and how far it usually misses."""
    as_of, diq, quarter, ref = context(latest)
    rows = latest.set_index("segment_name").reindex(list(TITLES))
    cells = SUMMARY_ROWS
    # Columns A to K: KPI, last comp, panel QTD YoY now, panel QTD YoY at the reference quarter,
    # panel change, slope, forecast, consensus, versus Street, model error, naive error.
    table = pd.DataFrame(
        {
            "KPI": rows.index.to_series().map(TITLES),
            f"Last reported comp ({ref})": rows["ref_actual"],
            f"Panel QTD YoY, {quarter} to day {diq}": rows["panel_yoy"],
            f"Panel QTD YoY, {ref} at day {diq}": rows["ref_panel_yoy"],
            "Panel change (pp)": [f"=100*(C{r}-D{r})" for r in cells],
            "Fitted slope": rows["coef"],
            f"Model forecast {quarter}": [f"=B{r}+F{r}*(C{r}-D{r})" for r in cells],
            "Consensus": rows["consensus"],
            "Model vs Street (pp)": [f"=100*(G{r}-H{r})" for r in cells],
            f"Typical model error at day {diq} (pp)": [
                mae_formula(quarters, title, f"Model (day {diq})")
                for title in TITLES.values()
            ],
            "Error if last comp is repeated (pp)": [
                mae_formula(quarters, title, "Naive") for title in TITLES.values()
            ],
        },
        index=rows.index,
    )
    table.to_excel(writer, sheet_name="Summary", index=False, startrow=2)
    ws = writer.sheets["Summary"]
    ws["A1"] = (
        f"URBN comparable-sales nowcast, as of {as_of:%Y-%m-%d} (day {diq} of fiscal {quarter})"
    )
    ws["A1"].font = Font(bold=True, size=13)
    formats = [number_format(column) for column in table.columns]
    style(ws, header_rows=range(3, 4), formats=formats, freeze_column=1)
    shade(
        ws,
        3,
        lambda label: "(pp)" in label and "error" not in label.lower(),
        diverging(4),
    )
    shade(
        ws, 3, lambda label: "(pp)" in label and "error" in label.lower(), sequential()
    )

    model_errors = f"J{cells.start}:J{cells[-1]}"
    lines = [
        "How to read this",
        (
            f"The model starts from the last reported comp ({ref}) and moves it by the fitted slope "
            "times the change in panel spend growth since that quarter, compared at the same day of "
            "the quarter: forecast = last comp + slope x (panel YoY now - panel YoY then)."
        ),
        (
            f'="Typical miss at day {diq}: "&TEXT(MIN({model_errors}),"0.0")&" to "'
            f'&TEXT(MAX({model_errors}),"0.0")&" pp, only a little better than repeating the last '
            f'comp. {DIRECTION}"'
        ),
        (
            "Reported comps are assumed usable 20 days after quarter end. The Notes sheet has "
            "definitions and assumptions."
        ),
    ]
    start = ws.max_row + 2
    for offset, line in enumerate(lines):
        ws.cell(row=start + offset, column=1, value=line)
    ws.cell(row=start, column=1).font = Font(bold=True)
    add_image(ws, charts["forecast_vs_actual"], f"A{start + len(lines) + 1}")
    add_image(ws, charts["error_by_day"], f"A{start + len(lines) + 1 + IMAGE_ROWS}")


def write_notes(writer, latest, quarters):
    """Who each sheet is for, definitions, the method in brief, assumptions and caveats."""
    as_of, diq, quarter, ref = context(latest)
    first_kpi = next(iter(TITLES.values()))
    reported = span(quarters, first_kpi, "Reported")
    model_errors = f"Summary!J{SUMMARY_ROWS.start}:J{SUMMARY_ROWS[-1]}"
    notes = [
        (
            "Data cut",
            (
                f"Card panel through {as_of:%Y-%m-%d}, day {diq} of fiscal {quarter}. Latest reported "
                f"quarter: {ref}. Consensus is the current snapshot for {quarter}; there is no "
                "consensus history."
            ),
        ),
        (
            "Summary",
            (
                "For the portfolio manager. Where the model sits for the in-progress quarter versus "
                "the last reported comp and consensus, what moved it, and how far a forecast made at "
                "this day has typically missed."
            ),
        ),
        (
            "By quarter",
            (
                f"For the PM and analyst. Every quarter's reported comp against the forecast made on "
                f"day {diq} and the naive baseline. The first chart on Summary, as numbers."
            ),
        ),
        (
            "By day",
            (
                "For the analyst. Mean absolute error of the model and the naive baseline for a "
                "forecast made on each day of the quarter, over the scored quarters. Shows how early "
                "the read can be trusted."
            ),
        ),
        (
            "Panel trends",
            (
                "For the analyst and PM. Daily, trailing-7-day and quarter-to-date panel spend growth "
                "year over year, aligned by day in quarter, for the last 13 months, newest first."
            ),
        ),
        (
            "Formulas",
            (
                "Panel change, the forecast, model vs Street and both error columns on Summary, the "
                "Error (pp) columns on By quarter and the averages on By day are live Excel formulas, "
                "so every derived number can be traced or overridden. Reported comps, panel YoY, "
                "consensus, fitted slopes, per-day errors and the trends are values from the pipeline."
            ),
        ),
        (
            "Comp",
            "Reported comparable-sales growth for the fiscal quarter, as a rate (6.2% = 0.062).",
        ),
        (
            "Panel QTD YoY",
            (
                "Quarter-to-date card-panel spend versus the same day in quarter of the prior fiscal "
                "year. The panel is a sample of URBN sales, not a comp; its coverage drifts, so the "
                "model uses changes in it rather than its level."
            ),
        ),
        (
            "pp",
            "Percentage points. A forecast of 6.2% against a reported 5.0% is an error of +1.2 pp.",
        ),
        (
            "Naive baseline",
            "Repeat the latest reported comp. The model has to beat this to add anything.",
        ),
        (
            "Model",
            (
                '="Forecast = last reported comp + slope x (panel QTD YoY this quarter minus panel QTD '
                'YoY of that reported quarter at the same day). "&"The slope is fitted per KPI and per '
                "day on earlier quarters only, and needs at least four of them, so the first scored "
                'quarter is "&\'By quarter\'!A3&"."'
            ),
        ),
        (
            "Availability",
            (
                "Reported comps are assumed usable 20 calendar days after the fiscal quarter ends; "
                "until then the reference is the quarter before. A simplifying assumption, not the "
                "actual release dates."
            ),
        ),
        (
            "Day in quarter",
            (
                "Year-over-year comparisons align day 1 to day 1 of the prior-year fiscal quarter. A "
                "day with no prior-year match (the leap-year tail) gets no YoY and no forecast."
            ),
        ),
        (
            "Confidence",
            (
                f'=COUNT({reported})&" scored quarters per KPI. Typical miss at day {diq} is "'
                f'&TEXT(MIN({model_errors}),"0.0")&" to "&TEXT(MAX({model_errors}),"0.0")'
                f'&" pp with only a small edge over the naive baseline. {DIRECTION}"'
            ),
        ),
    ]
    table = pd.DataFrame(notes, columns=["Topic", "Note"])
    table.to_excel(writer, sheet_name="Notes", index=False)
    ws = writer.sheets["Notes"]
    style(ws, header_rows=range(1, 2), formats=[None, None], freeze_column=1)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 18, 120
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def build(sheets, daily, charts, path):
    """Write the PM workbook from forecasting.py's sheets, the daily panel and the chart PNGs."""
    latest, backtest, errors = (
        sheets["Forecast"],
        sheets["Backtest"],
        sheets["Errors by day"],
    )
    _, diq, quarter, _ = context(latest)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        # By quarter first: Summary and Notes formulas point at its cells.
        quarters = write_by_quarter(writer, backtest, diq, quarter)
        write_by_day(writer, errors, diq)
        write_trends(writer, daily, charts["yoy_trends"])
        write_summary(writer, latest, charts, quarters)
        write_notes(writer, latest, quarters)
        writer.book.move_sheet("Summary", offset=-3)
        writer.book.active = 0
